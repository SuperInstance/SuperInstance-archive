"""
BI Export Engine
Export analytics data to various Business Intelligence tools and formats
"""

import pandas as pd
import numpy as np
from typing import Dict, List, Any, Optional, Union
from datetime import datetime, timedelta
from dataclasses import dataclass, field
from enum import Enum
import logging
import asyncio
import json
import csv
import io
import base64
from concurrent.futures import ThreadPoolExecutor

# Export format libraries
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from python_pptx import Presentation
    from python_pptx.util import Inches
    from python_pptx.enum.text import PP_ALIGN
    POWERPOINT_AVAILABLE = True
except ImportError:
    POWERPOINT_AVAILABLE = False

logger = logging.getLogger(__name__)

class ExportFormat(str, Enum):
    CSV = "csv"
    EXCEL = "excel"
    JSON = "json"
    PDF = "pdf"
    POWERPOINT = "powerpoint"
    PARQUET = "parquet"
    SQL_INSERT = "sql_insert"
    TABLEAU_EXTRACT = "tableau_extract"
    POWER_BI = "power_bi"

class BIToolType(str, Enum):
    TABLEAU = "tableau"
    POWER_BI = "power_bi"
    LOOKER = "looker"
    QLIK_SENSE = "qlik_sense"
    SUPERSET = "superset"
    METABASE = "metabase"
    GRAFANA = "grafana"

@dataclass
class ExportConfig:
    export_format: ExportFormat
    include_metadata: bool = True
    include_visualizations: bool = False
    date_format: str = "%Y-%m-%d %H:%M:%S"
    decimal_places: int = 4
    encoding: str = "utf-8"
    compression: Optional[str] = None  # gzip, zip, etc.
    custom_styling: Dict[str, Any] = field(default_factory=dict)
    bi_tool_settings: Dict[str, Any] = field(default_factory=dict)

@dataclass
class ExportResult:
    export_id: str
    export_format: str
    file_path: Optional[str]
    file_content: Optional[bytes]
    file_size: int
    export_time: float
    metadata: Dict[str, Any]
    success: bool
    error_message: Optional[str] = None

class BIExportEngine:
    """Engine for exporting analytics data to BI tools and formats"""
    
    def __init__(self):
        self.exports = {}
        self.executor = ThreadPoolExecutor(max_workers=4)
        logger.info("BI Export Engine initialized")
    
    async def export_data(self, data: Dict[str, Any], config: ExportConfig, filename: str = None) -> ExportResult:
        """Export data in the specified format"""
        start_time = datetime.utcnow()
        export_id = f"export_{int(start_time.timestamp())}_{config.export_format.value}"
        
        try:
            if filename is None:
                filename = f"analytics_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"
            
            # Convert data to appropriate format
            if config.export_format == ExportFormat.CSV:
                result = await self._export_csv(data, config, filename)
            elif config.export_format == ExportFormat.EXCEL:
                result = await self._export_excel(data, config, filename)
            elif config.export_format == ExportFormat.JSON:
                result = await self._export_json(data, config, filename)
            elif config.export_format == ExportFormat.PDF:
                result = await self._export_pdf(data, config, filename)
            elif config.export_format == ExportFormat.POWERPOINT:
                result = await self._export_powerpoint(data, config, filename)
            elif config.export_format == ExportFormat.PARQUET:
                result = await self._export_parquet(data, config, filename)
            elif config.export_format == ExportFormat.SQL_INSERT:
                result = await self._export_sql_insert(data, config, filename)
            elif config.export_format == ExportFormat.TABLEAU_EXTRACT:
                result = await self._export_tableau(data, config, filename)
            elif config.export_format == ExportFormat.POWER_BI:
                result = await self._export_power_bi(data, config, filename)
            else:
                raise ValueError(f"Unsupported export format: {config.export_format}")
            
            # Create export result
            export_result = ExportResult(
                export_id=export_id,
                export_format=config.export_format.value,
                file_path=result.get('file_path'),
                file_content=result.get('file_content'),
                file_size=result.get('file_size', 0),
                export_time=(datetime.utcnow() - start_time).total_seconds(),
                metadata={
                    'filename': filename,
                    'records_exported': result.get('records_exported', 0),
                    'columns_exported': result.get('columns_exported', 0),
                    'export_config': config.__dict__,
                    'created_at': start_time.isoformat()
                },
                success=True
            )
            
            # Store export information
            self.exports[export_id] = export_result
            
            logger.info(f"Data export completed successfully: {export_id}")
            return export_result
            
        except Exception as e:
            export_result = ExportResult(
                export_id=export_id,
                export_format=config.export_format.value,
                file_path=None,
                file_content=None,
                file_size=0,
                export_time=(datetime.utcnow() - start_time).total_seconds(),
                metadata={'error': str(e)},
                success=False,
                error_message=str(e)
            )
            
            self.exports[export_id] = export_result
            logger.error(f"Data export failed: {e}")
            return export_result
    
    async def export_analytics_results(self, job_id: str, job_results: Dict[str, Any], config: ExportConfig) -> ExportResult:
        """Export analytics job results"""
        export_data = {
            'job_metadata': {
                'job_id': job_id,
                'job_type': job_results.get('job_type'),
                'export_timestamp': datetime.utcnow().isoformat()
            },
            'results': job_results
        }
        
        filename = f"analytics_results_{job_id}_{config.export_format.value}"
        return await self.export_data(export_data, config, filename)
    
    async def create_bi_dashboard_export(self, dashboard_data: Dict[str, Any], bi_tool: BIToolType, config: ExportConfig) -> ExportResult:
        """Create export optimized for specific BI tools"""
        
        # Optimize data structure for specific BI tools
        optimized_data = await self._optimize_for_bi_tool(dashboard_data, bi_tool)
        
        # Set appropriate export format based on BI tool
        if bi_tool == BIToolType.TABLEAU:
            config.export_format = ExportFormat.TABLEAU_EXTRACT
        elif bi_tool == BIToolType.POWER_BI:
            config.export_format = ExportFormat.POWER_BI
        elif bi_tool in [BIToolType.LOOKER, BIToolType.SUPERSET, BIToolType.METABASE]:
            config.export_format = ExportFormat.CSV  # These tools work well with CSV
        elif bi_tool == BIToolType.GRAFANA:
            config.export_format = ExportFormat.JSON  # Grafana prefers JSON
        
        filename = f"bi_dashboard_{bi_tool.value}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"
        return await self.export_data(optimized_data, config, filename)
    
    async def get_export_info(self, export_id: str) -> Optional[ExportResult]:
        """Get information about an export"""
        return self.exports.get(export_id)
    
    async def list_exports(self, limit: int = 50) -> List[Dict[str, Any]]:
        """List recent exports"""
        exports_list = []
        for export_id, export_result in list(self.exports.items())[-limit:]:
            exports_list.append({
                'export_id': export_id,
                'export_format': export_result.export_format,
                'file_size': export_result.file_size,
                'export_time': export_result.export_time,
                'success': export_result.success,
                'created_at': export_result.metadata.get('created_at'),
                'filename': export_result.metadata.get('filename')
            })
        return exports_list
    
    # Format-specific export methods
    async def _export_csv(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data as CSV"""
        output = io.StringIO()
        
        # Convert data to DataFrame if needed
        df = await self._convert_to_dataframe(data)
        
        # Apply formatting
        df = await self._apply_formatting(df, config)
        
        # Write to CSV
        await asyncio.get_event_loop().run_in_executor(
            self.executor, 
            lambda: df.to_csv(output, index=False, encoding=config.encoding, date_format=config.date_format)
        )
        
        csv_content = output.getvalue().encode(config.encoding)
        file_size = len(csv_content)
        
        return {
            'file_content': csv_content,
            'file_size': file_size,
            'records_exported': len(df),
            'columns_exported': len(df.columns)
        }
    
    async def _export_excel(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data as Excel"""
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl not available. Please install: pip install openpyxl")
        
        output = io.BytesIO()
        
        # Convert data to DataFrame
        df = await self._convert_to_dataframe(data)
        df = await self._apply_formatting(df, config)
        
        # Create Excel file
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Analytics Data', index=False)
            
            # Apply custom styling if specified
            if config.custom_styling:
                await self._apply_excel_styling(writer.book, config.custom_styling)
            
            # Add metadata sheet if requested
            if config.include_metadata:
                metadata_df = pd.DataFrame([
                    {'Key': 'Export Date', 'Value': datetime.utcnow().isoformat()},
                    {'Key': 'Records', 'Value': len(df)},
                    {'Key': 'Columns', 'Value': len(df.columns)},
                    {'Key': 'Format', 'Value': config.export_format.value}
                ])
                metadata_df.to_excel(writer, sheet_name='Metadata', index=False)
        
        excel_content = output.getvalue()
        file_size = len(excel_content)
        
        return {
            'file_content': excel_content,
            'file_size': file_size,
            'records_exported': len(df),
            'columns_exported': len(df.columns)
        }
    
    async def _export_json(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data as JSON"""
        # Add metadata if requested
        if config.include_metadata:
            export_data = {
                'metadata': {
                    'export_date': datetime.utcnow().isoformat(),
                    'export_format': config.export_format.value,
                    'encoding': config.encoding
                },
                'data': data
            }
        else:
            export_data = data
        
        # Convert to JSON
        json_content = json.dumps(export_data, indent=2, default=str, ensure_ascii=False)
        json_bytes = json_content.encode(config.encoding)
        file_size = len(json_bytes)
        
        return {
            'file_content': json_bytes,
            'file_size': file_size,
            'records_exported': await self._count_records(data),
            'columns_exported': await self._count_columns(data)
        }
    
    async def _export_pdf(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data as PDF report"""
        if not PDF_AVAILABLE:
            raise ImportError("reportlab not available. Please install: pip install reportlab")
        
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=letter)
        elements = []
        
        # Styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=30,
            alignment=1  # Center alignment
        )
        
        # Title
        title = Paragraph("Analytics Report", title_style)
        elements.append(title)
        elements.append(Spacer(1, 20))
        
        # Convert data to DataFrame for table
        df = await self._convert_to_dataframe(data)
        df = await self._apply_formatting(df, config)
        
        # Limit rows for PDF (to prevent huge files)
        if len(df) > 100:
            df = df.head(100)
            note = Paragraph("Note: Only showing first 100 rows", styles['Italic'])
            elements.append(note)
            elements.append(Spacer(1, 10))
        
        # Create table data
        table_data = [df.columns.tolist()]
        for _, row in df.iterrows():
            table_data.append(row.tolist())
        
        # Create table
        table = Table(table_data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        elements.append(table)
        
        # Build PDF
        await asyncio.get_event_loop().run_in_executor(
            self.executor, doc.build, elements
        )
        
        pdf_content = output.getvalue()
        file_size = len(pdf_content)
        
        return {
            'file_content': pdf_content,
            'file_size': file_size,
            'records_exported': len(df),
            'columns_exported': len(df.columns)
        }
    
    async def _export_powerpoint(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data as PowerPoint presentation"""
        if not POWERPOINT_AVAILABLE:
            raise ImportError("python-pptx not available. Please install: pip install python-pptx")
        
        prs = Presentation()
        
        # Title slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = "Analytics Report"
        subtitle.text = f"Generated on {datetime.utcnow().strftime('%B %d, %Y')}"
        
        # Data slide
        bullet_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(bullet_slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        
        title.text = "Analytics Data Summary"
        
        # Create summary text
        df = await self._convert_to_dataframe(data)
        summary_text = f"• Total Records: {len(df)}\n"
        summary_text += f"• Total Columns: {len(df.columns)}\n"
        summary_text += f"• Export Date: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')}\n"
        
        content.text = summary_text
        
        # Save to bytes
        output = io.BytesIO()
        await asyncio.get_event_loop().run_in_executor(
            self.executor, prs.save, output
        )
        
        ppt_content = output.getvalue()
        file_size = len(ppt_content)
        
        return {
            'file_content': ppt_content,
            'file_size': file_size,
            'records_exported': len(df),
            'columns_exported': len(df.columns)
        }
    
    async def _export_parquet(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data as Parquet"""
        output = io.BytesIO()
        
        # Convert data to DataFrame
        df = await self._convert_to_dataframe(data)
        df = await self._apply_formatting(df, config)
        
        # Write to Parquet
        await asyncio.get_event_loop().run_in_executor(
            self.executor, df.to_parquet, output, index=False
        )
        
        parquet_content = output.getvalue()
        file_size = len(parquet_content)
        
        return {
            'file_content': parquet_content,
            'file_size': file_size,
            'records_exported': len(df),
            'columns_exported': len(df.columns)
        }
    
    async def _export_sql_insert(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data as SQL INSERT statements"""
        df = await self._convert_to_dataframe(data)
        df = await self._apply_formatting(df, config)
        
        table_name = config.bi_tool_settings.get('table_name', 'analytics_data')
        
        # Generate SQL INSERT statements
        sql_statements = []
        sql_statements.append(f"-- Analytics Data Export")
        sql_statements.append(f"-- Generated on {datetime.utcnow().isoformat()}")
        sql_statements.append(f"-- Records: {len(df)}")
        sql_statements.append("")
        
        # Create table statement
        create_table = f"CREATE TABLE IF NOT EXISTS {table_name} ("
        column_definitions = []
        for col in df.columns:
            if df[col].dtype in ['int64', 'int32']:
                col_type = "INTEGER"
            elif df[col].dtype in ['float64', 'float32']:
                col_type = "DECIMAL(18,4)"
            elif df[col].dtype == 'datetime64[ns]':
                col_type = "DATETIME"
            else:
                col_type = "VARCHAR(255)"
            column_definitions.append(f"    {col} {col_type}")
        
        create_table += ",\n".join(column_definitions) + "\n);"
        sql_statements.append(create_table)
        sql_statements.append("")
        
        # Insert statements
        for _, row in df.iterrows():
            values = []
            for val in row:
                if pd.isna(val):
                    values.append("NULL")
                elif isinstance(val, str):
                    values.append(f"'{val.replace(\"'\", \"''\")}'")
                else:
                    values.append(str(val))
            
            insert_statement = f"INSERT INTO {table_name} ({', '.join(df.columns)}) VALUES ({', '.join(values)});"
            sql_statements.append(insert_statement)
        
        sql_content = "\n".join(sql_statements)
        sql_bytes = sql_content.encode(config.encoding)
        file_size = len(sql_bytes)
        
        return {
            'file_content': sql_bytes,
            'file_size': file_size,
            'records_exported': len(df),
            'columns_exported': len(df.columns)
        }
    
    async def _export_tableau(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data optimized for Tableau"""
        # For Tableau, we'll export as CSV with specific formatting
        df = await self._convert_to_dataframe(data)
        
        # Tableau-specific formatting
        df.columns = [col.replace(' ', '_').replace('-', '_') for col in df.columns]
        
        # Ensure proper data types
        for col in df.columns:
            if df[col].dtype == 'object':
                try:
                    df[col] = pd.to_datetime(df[col])
                except:
                    pass
        
        output = io.StringIO()
        df.to_csv(output, index=False, encoding=config.encoding, date_format='%Y-%m-%d %H:%M:%S')
        
        csv_content = output.getvalue().encode(config.encoding)
        file_size = len(csv_content)
        
        return {
            'file_content': csv_content,
            'file_size': file_size,
            'records_exported': len(df),
            'columns_exported': len(df.columns)
        }
    
    async def _export_power_bi(self, data: Dict[str, Any], config: ExportConfig, filename: str) -> Dict[str, Any]:
        """Export data optimized for Power BI"""
        # Power BI works well with Excel files
        return await self._export_excel(data, config, filename)
    
    # Helper methods
    async def _convert_to_dataframe(self, data: Dict[str, Any]) -> pd.DataFrame:
        """Convert various data formats to DataFrame"""
        if isinstance(data, pd.DataFrame):
            return data
        elif isinstance(data, dict):
            if 'results' in data and isinstance(data['results'], dict):
                # Handle nested results structure
                results_data = data['results']
                if any(isinstance(v, (list, dict)) for v in results_data.values()):
                    # Try to flatten the structure
                    flattened_data = {}
                    for key, value in results_data.items():
                        if isinstance(value, list):
                            flattened_data[key] = value
                        elif isinstance(value, dict):
                            for subkey, subvalue in value.items():
                                flattened_data[f"{key}_{subkey}"] = [subvalue] if not isinstance(subvalue, list) else subvalue
                        else:
                            flattened_data[key] = [value]
                    
                    # Make all lists the same length
                    max_length = max(len(v) if isinstance(v, list) else 1 for v in flattened_data.values())
                    for key, value in flattened_data.items():
                        if isinstance(value, list) and len(value) < max_length:
                            flattened_data[key] = value + [None] * (max_length - len(value))
                        elif not isinstance(value, list):
                            flattened_data[key] = [value] * max_length
                    
                    return pd.DataFrame(flattened_data)
                else:
                    return pd.DataFrame([results_data])
            else:
                return pd.DataFrame([data])
        elif isinstance(data, list):
            return pd.DataFrame(data)
        else:
            return pd.DataFrame([{'value': data}])
    
    async def _apply_formatting(self, df: pd.DataFrame, config: ExportConfig) -> pd.DataFrame:
        """Apply formatting to DataFrame"""
        formatted_df = df.copy()
        
        # Format numeric columns
        for col in formatted_df.select_dtypes(include=[np.number]).columns:
            if formatted_df[col].dtype in ['float64', 'float32']:
                formatted_df[col] = formatted_df[col].round(config.decimal_places)
        
        # Format datetime columns
        for col in formatted_df.select_dtypes(include=['datetime64']).columns:
            formatted_df[col] = formatted_df[col].dt.strftime(config.date_format)
        
        return formatted_df
    
    async def _apply_excel_styling(self, workbook, styling: Dict[str, Any]):
        """Apply custom styling to Excel workbook"""
        # This is a placeholder for custom Excel styling
        # In a real implementation, you would apply fonts, colors, etc.
        pass
    
    async def _optimize_for_bi_tool(self, data: Dict[str, Any], bi_tool: BIToolType) -> Dict[str, Any]:
        """Optimize data structure for specific BI tools"""
        if bi_tool == BIToolType.TABLEAU:
            # Tableau prefers flat structures with proper column names
            return await self._flatten_data_for_tableau(data)
        elif bi_tool == BIToolType.POWER_BI:
            # Power BI works well with relational structures
            return await self._structure_for_power_bi(data)
        else:
            return data
    
    async def _flatten_data_for_tableau(self, data: Dict[str, Any]) -> Dict[str, Any]:
        """Flatten data structure for Tableau"""
        # Implement Tableau-specific data flattening
        return data
    
    async def _structure_for_power_bi(self, data: Dict[str, Any]) -> Dict[str, Any]:
        """Structure data for Power BI"""
        # Implement Power BI-specific data structuring
        return data
    
    async def _count_records(self, data: Any) -> int:
        """Count records in data"""
        if isinstance(data, pd.DataFrame):
            return len(data)
        elif isinstance(data, list):
            return len(data)
        elif isinstance(data, dict):
            if 'results' in data:
                return await self._count_records(data['results'])
            else:
                return 1
        else:
            return 1
    
    async def _count_columns(self, data: Any) -> int:
        """Count columns in data"""
        if isinstance(data, pd.DataFrame):
            return len(data.columns)
        elif isinstance(data, dict):
            return len(data.keys())
        elif isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
            return len(data[0].keys())
        else:
            return 1