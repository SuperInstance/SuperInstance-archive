"""
Document diff and comparison engine
"""

import asyncio
import difflib
import json
import logging
import mimetypes
import re
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional

# Document processing libraries
try:
    import docx
    from docx.document import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import PyPDF2
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    from openpyxl import load_workbook
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False

logger = logging.getLogger(__name__)

class DiffEngine:
    def __init__(self):
        self.supported_formats = {
            'text': ['.txt', '.md', '.py', '.js', '.html', '.css', '.json', '.xml', '.csv'],
            'docx': ['.docx'] if HAS_DOCX else [],
            'pdf': ['.pdf'] if HAS_PDF else [],
            'excel': ['.xlsx', '.xls'] if HAS_EXCEL else [],
            'binary': ['.jpg', '.jpeg', '.png', '.gif', '.zip', '.rar', '.tar', '.gz']
        }
        
    async def initialize(self):
        """Initialize diff engine"""
        logger.info("Initializing diff engine")
        
        # Log available document processors
        processors = []
        if HAS_DOCX:
            processors.append("DOCX")
        if HAS_PDF:
            processors.append("PDF")
        if HAS_EXCEL:
            processors.append("Excel")
        
        logger.info(f"Available document processors: {', '.join(processors) if processors else 'Text only'}")
        
    async def cleanup(self):
        """Cleanup resources"""
        logger.info("Cleaning up diff engine")
        
    async def generate_diff(self, file_path_1: str, file_path_2: str, 
                          file_type: str = None) -> Dict[str, Any]:
        """Generate diff between two files"""
        try:
            path_1 = Path(file_path_1)
            path_2 = Path(file_path_2)
            
            if not path_1.exists() or not path_2.exists():
                raise FileNotFoundError("One or both files do not exist")
            
            # Determine file type if not provided
            if not file_type:
                file_type = self._detect_file_type(path_1)
            
            # Extract content based on file type
            content_1 = await self._extract_content(path_1, file_type)
            content_2 = await self._extract_content(path_2, file_type)
            
            # Generate diff
            diff_result = self._generate_text_diff(content_1, content_2)
            
            # Add file metadata
            diff_result.update({
                "file_1": {
                    "path": str(path_1),
                    "size": path_1.stat().st_size,
                    "modified": path_1.stat().st_mtime
                },
                "file_2": {
                    "path": str(path_2),
                    "size": path_2.stat().st_size,
                    "modified": path_2.stat().st_mtime
                },
                "file_type": file_type
            })
            
            return diff_result
            
        except Exception as e:
            logger.error(f"Failed to generate diff between {file_path_1} and {file_path_2}: {e}")
            raise
            
    async def compare_files(self, file_path_1: str, file_path_2: str, 
                          file_type: str = None) -> Dict[str, Any]:
        """Compare two files and return detailed analysis"""
        try:
            diff_result = await self.generate_diff(file_path_1, file_path_2, file_type)
            
            # Add comparison summary
            comparison = {
                "identical": diff_result["changes_count"] == 0,
                "similarity_ratio": diff_result["similarity_ratio"],
                "size_difference": diff_result["file_2"]["size"] - diff_result["file_1"]["size"],
                "lines_added": diff_result["lines_added"],
                "lines_removed": diff_result["lines_removed"],
                "lines_modified": diff_result["lines_modified"],
                "diff_html": diff_result["diff_html"],
                "summary": diff_result["summary"]
            }
            
            # Add file-type specific analysis
            if file_type in ["docx", "pdf", "excel"]:
                comparison["structured_changes"] = await self._analyze_structured_changes(
                    file_path_1, file_path_2, file_type
                )
            
            comparison.update(diff_result)
            return comparison
            
        except Exception as e:
            logger.error(f"Failed to compare files {file_path_1} and {file_path_2}: {e}")
            raise
            
    def _detect_file_type(self, file_path: Path) -> str:
        """Detect file type from extension"""
        extension = file_path.suffix.lower()
        
        for format_type, extensions in self.supported_formats.items():
            if extension in extensions:
                return format_type
                
        return "binary"  # Default to binary for unknown types
        
    async def _extract_content(self, file_path: Path, file_type: str) -> str:
        """Extract text content from file based on type"""
        try:
            if file_type == "text":
                return await self._extract_text_content(file_path)
            elif file_type == "docx" and HAS_DOCX:
                return await self._extract_docx_content(file_path)
            elif file_type == "pdf" and HAS_PDF:
                return await self._extract_pdf_content(file_path)
            elif file_type == "excel" and HAS_EXCEL:
                return await self._extract_excel_content(file_path)
            elif file_type == "binary":
                return f"[Binary file: {file_path.name}, size: {file_path.stat().st_size} bytes]"
            else:
                # Fallback to text extraction
                return await self._extract_text_content(file_path)
                
        except Exception as e:
            logger.warning(f"Failed to extract content from {file_path}: {e}")
            return f"[Content extraction failed: {str(e)}]"
            
    async def _extract_text_content(self, file_path: Path) -> str:
        """Extract content from text files"""
        try:
            # Try different encodings
            encodings = ['utf-8', 'utf-16', 'iso-8859-1', 'cp1252']
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        return f.read()
                except UnicodeDecodeError:
                    continue
                    
            # If all encodings fail, read as binary and decode with errors='ignore'
            with open(file_path, 'rb') as f:
                return f.read().decode('utf-8', errors='ignore')
                
        except Exception as e:
            raise Exception(f"Failed to read text file: {e}")
            
    async def _extract_docx_content(self, file_path: Path) -> str:
        """Extract text content from DOCX files"""
        try:
            doc = docx.Document(str(file_path))
            paragraphs = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    paragraphs.append(paragraph.text)
                    
            # Extract table content
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        paragraphs.append(" | ".join(row_text))
            
            return "\n".join(paragraphs)
            
        except Exception as e:
            raise Exception(f"Failed to extract DOCX content: {e}")
            
    async def _extract_pdf_content(self, file_path: Path) -> str:
        """Extract text content from PDF files"""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text_content = []
                
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    page_text = page.extract_text()
                    if page_text.strip():
                        text_content.append(f"--- Page {page_num + 1} ---")
                        text_content.append(page_text)
                
                return "\n".join(text_content)
                
        except Exception as e:
            raise Exception(f"Failed to extract PDF content: {e}")
            
    async def _extract_excel_content(self, file_path: Path) -> str:
        """Extract content from Excel files"""
        try:
            workbook = load_workbook(str(file_path), read_only=True)
            content_parts = []
            
            for sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                content_parts.append(f"--- Sheet: {sheet_name} ---")
                
                for row in worksheet.iter_rows(values_only=True):
                    row_values = [str(cell) if cell is not None else "" for cell in row]
                    if any(val.strip() for val in row_values):  # Skip empty rows
                        content_parts.append(" | ".join(row_values))
                        
            workbook.close()
            return "\n".join(content_parts)
            
        except Exception as e:
            raise Exception(f"Failed to extract Excel content: {e}")
            
    def _generate_text_diff(self, content_1: str, content_2: str) -> Dict[str, Any]:
        """Generate detailed text diff"""
        try:
            lines_1 = content_1.splitlines(keepends=True)
            lines_2 = content_2.splitlines(keepends=True)
            
            # Calculate similarity ratio
            similarity = difflib.SequenceMatcher(None, content_1, content_2).ratio()
            
            # Generate unified diff
            unified_diff = list(difflib.unified_diff(
                lines_1, lines_2,
                fromfile="Version 1", tofile="Version 2",
                lineterm=""
            ))
            
            # Generate HTML diff
            html_diff = difflib.HtmlDiff()
            diff_html = html_diff.make_file(
                lines_1, lines_2,
                fromdesc="Version 1", todesc="Version 2"
            )
            
            # Count changes
            changes = self._count_changes(unified_diff)
            
            # Generate summary
            summary = self._generate_diff_summary(changes, similarity)
            
            return {
                "unified_diff": unified_diff,
                "diff_html": diff_html,
                "similarity_ratio": similarity,
                "changes_count": changes["total"],
                "lines_added": changes["added"],
                "lines_removed": changes["removed"],
                "lines_modified": changes["modified"],
                "summary": summary
            }
            
        except Exception as e:
            logger.error(f"Failed to generate text diff: {e}")
            raise
            
    def _count_changes(self, unified_diff: List[str]) -> Dict[str, int]:
        """Count different types of changes in unified diff"""
        changes = {
            "added": 0,
            "removed": 0,
            "modified": 0,
            "total": 0
        }
        
        for line in unified_diff:
            if line.startswith('+') and not line.startswith('+++'):
                changes["added"] += 1
            elif line.startswith('-') and not line.startswith('---'):
                changes["removed"] += 1
                
        # Estimate modified lines (lines that are both added and removed)
        changes["modified"] = min(changes["added"], changes["removed"])
        changes["added"] -= changes["modified"]
        changes["removed"] -= changes["modified"]
        changes["total"] = changes["added"] + changes["removed"] + changes["modified"]
        
        return changes
        
    def _generate_diff_summary(self, changes: Dict[str, int], similarity: float) -> str:
        """Generate human-readable diff summary"""
        if changes["total"] == 0:
            return "Files are identical"
        
        summary_parts = []
        
        if changes["added"] > 0:
            summary_parts.append(f"{changes['added']} line{'s' if changes['added'] != 1 else ''} added")
        
        if changes["removed"] > 0:
            summary_parts.append(f"{changes['removed']} line{'s' if changes['removed'] != 1 else ''} removed")
        
        if changes["modified"] > 0:
            summary_parts.append(f"{changes['modified']} line{'s' if changes['modified'] != 1 else ''} modified")
        
        summary = ", ".join(summary_parts)
        similarity_percent = int(similarity * 100)
        summary += f" ({similarity_percent}% similar)"
        
        return summary
        
    async def _analyze_structured_changes(self, file_path_1: str, file_path_2: str, 
                                        file_type: str) -> Dict[str, Any]:
        """Analyze changes in structured documents"""
        try:
            if file_type == "docx" and HAS_DOCX:
                return await self._analyze_docx_changes(file_path_1, file_path_2)
            elif file_type == "excel" and HAS_EXCEL:
                return await self._analyze_excel_changes(file_path_1, file_path_2)
            else:
                return {}
                
        except Exception as e:
            logger.warning(f"Failed to analyze structured changes: {e}")
            return {}
            
    async def _analyze_docx_changes(self, file_path_1: str, file_path_2: str) -> Dict[str, Any]:
        """Analyze changes specific to DOCX documents"""
        try:
            doc1 = docx.Document(file_path_1)
            doc2 = docx.Document(file_path_2)
            
            analysis = {
                "paragraph_changes": {
                    "added": max(0, len(doc2.paragraphs) - len(doc1.paragraphs)),
                    "removed": max(0, len(doc1.paragraphs) - len(doc2.paragraphs))
                },
                "table_changes": {
                    "added": max(0, len(doc2.tables) - len(doc1.tables)),
                    "removed": max(0, len(doc1.tables) - len(doc2.tables))
                }
            }
            
            return analysis
            
        except Exception as e:
            logger.warning(f"Failed to analyze DOCX changes: {e}")
            return {}
            
    async def _analyze_excel_changes(self, file_path_1: str, file_path_2: str) -> Dict[str, Any]:
        """Analyze changes specific to Excel documents"""
        try:
            wb1 = load_workbook(file_path_1, read_only=True)
            wb2 = load_workbook(file_path_2, read_only=True)
            
            analysis = {
                "worksheet_changes": {
                    "added": len(set(wb2.sheetnames) - set(wb1.sheetnames)),
                    "removed": len(set(wb1.sheetnames) - set(wb2.sheetnames)),
                    "common": len(set(wb1.sheetnames) & set(wb2.sheetnames))
                }
            }
            
            # Analyze common worksheets
            for sheet_name in set(wb1.sheetnames) & set(wb2.sheetnames):
                ws1 = wb1[sheet_name]
                ws2 = wb2[sheet_name]
                
                max_row_1 = ws1.max_row
                max_row_2 = ws2.max_row
                max_col_1 = ws1.max_column
                max_col_2 = ws2.max_column
                
                analysis[f"sheet_{sheet_name}"] = {
                    "row_changes": max_row_2 - max_row_1,
                    "column_changes": max_col_2 - max_col_1
                }
            
            wb1.close()
            wb2.close()
            
            return analysis
            
        except Exception as e:
            logger.warning(f"Failed to analyze Excel changes: {e}")
            return {}
            
    async def create_patch(self, original_file: str, modified_file: str) -> str:
        """Create a patch file for the differences"""
        try:
            diff_result = await self.generate_diff(original_file, modified_file)
            
            # Convert unified diff to patch format
            patch_content = "\n".join(diff_result["unified_diff"])
            
            return patch_content
            
        except Exception as e:
            logger.error(f"Failed to create patch: {e}")
            raise
            
    def get_supported_formats(self) -> Dict[str, List[str]]:
        """Get list of supported file formats"""
        return self.supported_formats.copy()